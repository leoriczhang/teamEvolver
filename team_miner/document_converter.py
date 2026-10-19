"""Convert knowledge workspace uploads to UTF-8 Markdown before OpenViking import."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree

SUPPORTED_KNOWLEDGE_SUFFIXES = {".md", ".markdown", ".txt", ".docx", ".xlsx", ".pdf"}
SUPPORTED_UPLOAD_SUFFIXES = SUPPORTED_KNOWLEDGE_SUFFIXES | {".csv", ".pptx"}
MAX_ARCHIVE_FILES = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_NORMALIZED_DOCUMENT_CHARS = 5_000_000
MAX_SPREADSHEET_CELLS = 200_000
MAX_PDF_PAGES = 2_000


@dataclass(frozen=True)
class ConvertedKnowledgeUpload:
    markdown: str
    source_format: str
    source_encoding: str


def _decode_text(raw: bytes) -> tuple[str, str]:
    if not raw:
        raise ValueError("不允许上传空文件")
    if b"\x00" in raw:
        raise ValueError("文件包含二进制内容，请检查文件格式")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文件编码，请转换为 UTF-8 文本后上传")


def _finish_markdown(text: str, filename: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip("\ufeff\n")
    if not re.sub(r"\s+", "", normalized):
        raise ValueError(f"{filename} 转换后没有可用文本")
    if len(normalized) > MAX_NORMALIZED_DOCUMENT_CHARS:
        raise ValueError(f"{filename} 转换后的文本超过 500 万字符，请拆分后上传")
    return normalized + "\n"


def _validate_office_archive(raw: bytes, filename: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_FILES:
                raise ValueError(f"{filename} 内部文件数量过多")
            if sum(entry.file_size for entry in entries) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ValueError(f"{filename} 解压后超过 100 MB")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValueError(f"{filename} 已加密，暂不支持解析")
    except ValueError:
        raise
    except zipfile.BadZipFile as exc:
        raise ValueError(f"{filename} 不是有效的 Office 文件") from exc


def _markdown_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\r\n", "<br>").replace("\r", "<br>").replace("\n", "<br>").strip()


def _markdown_table(rows: list[list[str]]) -> str:
    cleaned: list[list[str]] = []
    width = 0
    for row in rows:
        cells = [_markdown_cell(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        cleaned.append(cells)
        width = max(width, len(cells))
    if not cleaned or not width:
        return ""
    padded = [row + [""] * (width - len(row)) for row in cleaned]
    header = [value or f"列 {index + 1}" for index, value in enumerate(padded[0])]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in padded[1:])
    return "\n".join(lines)


def _docx_to_markdown(raw: bytes, filename: str) -> str:
    _validate_office_archive(raw, filename)
    try:
        from docx import Document

        document = Document(io.BytesIO(raw))
    except Exception as exc:
        raise ValueError(f"{filename} 不是有效的 DOCX 文件") from exc
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = str(getattr(paragraph.style, "name", "") or "").lower()
        match = re.match(r"heading\s+([1-6])", style)
        parts.append(f"{'#' * int(match.group(1))} {text}" if match else text)
    for table in document.tables:
        rendered = _markdown_table(
            [[str(cell.text or "") for cell in row.cells] for row in table.rows]
        )
        if rendered:
            parts.append(rendered)
    return _finish_markdown("\n\n".join(parts), filename)


def _xlsx_to_markdown(raw: bytes, filename: str) -> str:
    _validate_office_archive(raw, filename)
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"{filename} 不是有效的 XLSX 文件") from exc
    parts: list[str] = []
    cell_count = 0
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                cell_count += len(values)
                if cell_count > MAX_SPREADSHEET_CELLS:
                    raise ValueError(f"{filename} 超过 20 万个单元格，请拆分后上传")
                rows.append(values)
            table = _markdown_table(rows)
            if table:
                parts.append(f"## {sheet.title}\n\n{table}")
    finally:
        workbook.close()
    title = Path(filename).stem.replace("\n", " ").strip() or "知识文档"
    return _finish_markdown(f"# {title}\n\n" + "\n\n".join(parts), filename)


def _pdf_to_markdown(raw: bytes, filename: str) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
    except Exception as exc:
        raise ValueError(f"{filename} 不是有效的 PDF 文件") from exc
    if reader.is_encrypted:
        raise ValueError(f"{filename} 已加密，暂不支持解析")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"{filename} 超过 2000 页，请拆分后上传")
    title = Path(filename).stem.replace("\n", " ").strip() or "知识文档"
    parts = [f"# {title}"]
    for index, page in enumerate(reader.pages, start=1):
        text = str(page.extract_text() or "").strip()
        if text:
            parts.append(f"## 第 {index} 页\n\n{text}")
    return _finish_markdown("\n\n".join(parts), filename)


def normalize_knowledge_document(filename: str, raw: bytes) -> ConvertedKnowledgeUpload:
    """Convert Markdown, text, DOCX, XLSX, or PDF into UTF-8 Markdown."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_KNOWLEDGE_SUFFIXES:
        allowed = "、".join(sorted(SUPPORTED_KNOWLEDGE_SUFFIXES))
        raise ValueError(f"不支持 {suffix or '无扩展名'} 文件；支持：{allowed}")
    if suffix in {".md", ".markdown"}:
        text, encoding = _decode_text(raw)
        return ConvertedKnowledgeUpload(_finish_markdown(text, filename), "markdown", encoding)
    if suffix == ".txt":
        text, encoding = _decode_text(raw)
        title = Path(filename).stem.replace("\n", " ").strip() or "知识文档"
        return ConvertedKnowledgeUpload(
            _finish_markdown(f"# {title}\n\n{text}", filename),
            "txt",
            encoding,
        )
    if suffix == ".docx":
        return ConvertedKnowledgeUpload(_docx_to_markdown(raw, filename), "docx", "binary")
    if suffix == ".xlsx":
        return ConvertedKnowledgeUpload(_xlsx_to_markdown(raw, filename), "xlsx", "binary")
    return ConvertedKnowledgeUpload(_pdf_to_markdown(raw, filename), "pdf", "binary")


def _csv_to_markdown(raw: bytes, filename: str) -> tuple[str, str]:
    text, encoding = _decode_text(raw)
    try:
        rows = list(csv.reader(io.StringIO(text)))
    except csv.Error as exc:
        raise ValueError(f"{filename} 不是有效的 CSV 文件") from exc
    if sum(len(row) for row in rows) > MAX_SPREADSHEET_CELLS:
        raise ValueError(f"{filename} 超过 20 万个单元格，请拆分后上传")
    table = _markdown_table(rows)
    if not table:
        raise ValueError(f"{filename} 转换后没有可用文本")
    title = Path(filename).stem.replace("\n", " ").strip() or "知识文档"
    return _finish_markdown(f"# {title}\n\n{table}", filename), encoding


def _pptx_to_markdown(raw: bytes, filename: str) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_FILES:
                raise ValueError(f"{filename} 内部文件数量过多")
            if sum(entry.file_size for entry in entries) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ValueError(f"{filename} 解压后超过 100 MB")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValueError(f"{filename} 已加密，暂不支持解析")
            slide_names = sorted(
                (
                    entry.filename
                    for entry in entries
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", entry.filename)
                ),
                key=lambda value: int(re.search(r"(\d+)", Path(value).stem).group(1)),
            )
            slides: list[list[str]] = []
            for slide_name in slide_names:
                root = ElementTree.fromstring(archive.read(slide_name))
                slides.append([
                    str(node.text or "").strip()
                    for node in root.iter()
                    if node.tag.endswith("}t") and str(node.text or "").strip()
                ])
    except ValueError:
        raise
    except (zipfile.BadZipFile, ElementTree.ParseError, AttributeError) as exc:
        raise ValueError(f"{filename} 不是有效的 PPTX 文件") from exc

    title = Path(filename).stem.replace("\n", " ").strip() or "知识文档"
    parts = [f"# {title}"]
    for index, texts in enumerate(slides, start=1):
        if texts:
            parts.append(f"## 第 {index} 页\n\n" + "\n\n".join(texts))
    if len(parts) == 1:
        raise ValueError(f"{filename} 转换后没有可用文本")
    return _finish_markdown("\n\n".join(parts), filename)


def convert_knowledge_upload(filename: str, raw: bytes) -> ConvertedKnowledgeUpload:
    """Normalize a supported browser upload to one Markdown document."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_UPLOAD_SUFFIXES:
        allowed = "、".join(sorted(SUPPORTED_UPLOAD_SUFFIXES))
        raise ValueError(f"不支持 {suffix or '无扩展名'} 文件；支持：{allowed}")

    if suffix == ".csv":
        markdown, encoding = _csv_to_markdown(raw, filename)
        converted = ConvertedKnowledgeUpload(markdown, "csv", encoding)
    elif suffix == ".pptx":
        converted = ConvertedKnowledgeUpload(_pptx_to_markdown(raw, filename), "pptx", "binary")
    else:
        normalized = normalize_knowledge_document(filename, raw)
        converted = ConvertedKnowledgeUpload(
            normalized.markdown,
            normalized.source_format,
            normalized.source_encoding,
        )

    if suffix in {".md", ".markdown"}:
        return converted
    metadata = (
        "---\n"
        f"source_file: {json.dumps(Path(filename).name, ensure_ascii=False)}\n"
        f"source_format: {converted.source_format}\n"
        f"converted_at: {datetime.now(timezone.utc).isoformat()}\n"
        "---\n\n"
    )
    return ConvertedKnowledgeUpload(
        metadata + converted.markdown,
        converted.source_format,
        converted.source_encoding,
    )
