from __future__ import annotations

import io
import zipfile

import pytest

from team_miner.document_converter import convert_knowledge_upload


def test_markdown_upload_normalizes_legacy_encoding_without_adding_frontmatter():
    converted = convert_knowledge_upload("guide.md", "新的领域知识".encode("gb18030"))
    assert converted.markdown == "新的领域知识\n"
    assert converted.source_format == "md"
    assert converted.source_encoding == "gb18030"


def test_plain_text_upload_becomes_markdown_with_source_metadata():
    converted = convert_knowledge_upload("规则.txt", "时限是 7 天".encode())
    assert converted.markdown.startswith('---\nsource_file: "规则.txt"\nsource_format: txt\n')
    assert "时限是 7 天" in converted.markdown
    assert converted.source_format == "txt"


def test_docx_upload_preserves_headings_paragraphs_and_tables():
    from docx import Document

    document = Document()
    document.add_heading("退款规则", level=1)
    document.add_paragraph("签收后七天内可以申请退款。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "场景"
    table.cell(0, 1).text = "时限"
    table.cell(1, 0).text = "退款"
    table.cell(1, 1).text = "7 天"
    payload = io.BytesIO()
    document.save(payload)

    converted = convert_knowledge_upload("售后手册.docx", payload.getvalue())
    assert 'source_file: "售后手册.docx"' in converted.markdown
    assert "# 售后手册\n" in converted.markdown
    assert "## 退款规则" in converted.markdown
    assert "签收后七天内可以申请退款。" in converted.markdown
    assert "| 场景 | 时限 |" in converted.markdown
    assert "| 退款 | 7 天 |" in converted.markdown
    assert converted.source_format == "docx"


def test_xlsx_upload_preserves_all_sheets_and_formulas():
    from openpyxl import Workbook

    workbook = Workbook()
    rules = workbook.active
    rules.title = "规则"
    rules.append(["场景", "时限"])
    rules.append(["退款", 7])
    formulas = workbook.create_sheet("计算")
    formulas.append(["项目", "值"])
    formulas.append(["合计", "=SUM(1,2)"])
    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()

    converted = convert_knowledge_upload("规则表.xlsx", payload.getvalue())
    assert "# 规则表\n" in converted.markdown
    assert "## 工作表：规则" in converted.markdown
    assert "| 退款 | 7 |" in converted.markdown
    assert "## 工作表：计算" in converted.markdown
    assert "| 合计 | =SUM(1,2) |" in converted.markdown
    assert converted.source_format == "xlsx"


def _pdf(*, with_text: bool) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    if with_text:
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 100 Td (Refund within 7 days) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
    payload = io.BytesIO()
    writer.write(payload)
    return payload.getvalue()


def test_text_pdf_upload_extracts_text_and_scanned_pdf_requires_ocr():
    converted = convert_knowledge_upload("policy.pdf", _pdf(with_text=True))
    assert "## 第 1 页" in converted.markdown
    assert "Refund within 7 days" in converted.markdown
    assert converted.source_format == "pdf"
    with pytest.raises(ValueError, match="OCR"):
        convert_knowledge_upload("scan.pdf", _pdf(with_text=False))


def test_csv_upload_escapes_cells_and_preserves_unicode():
    converted = convert_knowledge_upload("规则.csv", '场景,规则\n退款,"7 天 | 特例"\n'.encode())
    assert "| 场景 | 规则 |" in converted.markdown
    assert "| 退款 | 7 天 \\| 特例 |" in converted.markdown
    assert converted.source_format == "csv"


def test_pptx_upload_extracts_slide_text_in_numeric_order():
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        for number in (10, 2, 1):
            archive.writestr(
                f"ppt/slides/slide{number}.xml",
                f'<slide xmlns:a="urn:drawing"><a:t>Evidence {number}</a:t></slide>',
            )
    converted = convert_knowledge_upload("说明.pptx", payload.getvalue())
    assert converted.markdown.index("Evidence 1\n") < converted.markdown.index("Evidence 2\n")
    assert converted.markdown.index("Evidence 2\n") < converted.markdown.index("Evidence 10\n")
    assert converted.source_format == "pptx"


@pytest.mark.parametrize("filename,content", [
    ("empty.md", b""), ("binary.txt", b"a\x00b"), ("image.png", b"image"),
    ("invalid.pdf", b"not a pdf"), ("invalid.pptx", b"not a zip"), ("empty.csv", b"\n\n"),
])
def test_invalid_or_empty_upload_is_rejected(filename, content):
    with pytest.raises(ValueError):
        convert_knowledge_upload(filename, content)


@pytest.mark.parametrize("extension", ["docx", "xlsx"])
def test_empty_office_upload_is_rejected(extension):
    payload = io.BytesIO()
    if extension == "docx":
        from docx import Document

        Document().save(payload)
    else:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.save(payload)
        workbook.close()
    with pytest.raises(ValueError, match="可用文本"):
        convert_knowledge_upload(f"empty.{extension}", payload.getvalue())
