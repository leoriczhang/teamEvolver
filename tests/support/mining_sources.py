"""In-memory OpenViking transport for source-management workflow contracts."""

from __future__ import annotations

import base64
import io
import zipfile

from starlette.exceptions import HTTPException

from team_miner.skill_miner.config import MiningSettings
from team_miner.skill_miner.engine_resources import MiningResourceMixin


class SourceWorkspaceRemote(MiningResourceMixin):
    def __init__(self, repository, root, *, account="team"):
        self.repository = repository
        self.settings = MiningSettings(1, 0, {}, {"account_id": account})
        self.task = None
        self.account = account
        self.directories = {"viking://resources", root}
        self.files = {}
        self.calls = []
        self.unavailable_roots = set()

    def _check_account(self, account):
        if account != self.account:
            raise HTTPException(403, "TENANT_ACCOUNT_MISMATCH")

    async def _knowledge_request(self, account, method, path, **kwargs):
        self._check_account(account)
        self.calls.append((method, path, kwargs))
        params, body = kwargs.get("params", {}), kwargs.get("json", {})
        uri = str(params.get("uri") or body.get("uri") or "")
        if path == "/api/v1/fs/ls":
            if uri in self.unavailable_roots:
                raise HTTPException(503, "upstream temporarily unavailable")
            if uri not in self.directories:
                raise HTTPException(404, "NOT_FOUND")
            rows = [{"uri": child, "isDir": True} for child in self.directories
                    if child.rsplit("/", 1)[0] == uri]
            rows += [{"uri": child, "isDir": False, "size": len(content.encode())}
                     for child, content in self.files.items() if child.rsplit("/", 1)[0] == uri]
            rows.sort(key=lambda row: (not row["isDir"], row["uri"].rsplit("/", 1)[-1].lower(),
                                       row["uri"].rsplit("/", 1)[-1]))
            offset = params.get("offset", 0)
            return rows[offset:offset + params.get("node_limit", 1000)]
        if path == "/api/v1/fs/mkdir":
            if uri in self.directories:
                raise HTTPException(409, "ALREADY_EXISTS")
            if uri.rsplit("/", 1)[0] not in self.directories:
                raise HTTPException(404, "NOT_FOUND")
            self.directories.add(uri)
            return {"uri": uri}
        if path == "/api/v1/content/read":
            if uri not in self.files:
                raise HTTPException(404, "NOT_FOUND")
            return {"content": self.files[uri]}
        if path == "/api/v1/content/write":
            if body["mode"] == "create" and uri in self.files:
                raise HTTPException(409, "ALREADY_EXISTS")
            if uri.rsplit("/", 1)[0] not in self.directories:
                raise HTTPException(404, "NOT_FOUND")
            self.files[uri] = body["content"]
            return {"uri": uri}
        if method == "DELETE" and path == "/api/v1/fs":
            self.files = {key: value for key, value in self.files.items()
                          if key != uri and not key.startswith(uri + "/")}
            self.directories = {key for key in self.directories if key != uri and not key.startswith(uri + "/")}
            return {"uri": uri}
        raise AssertionError((method, path))


def source_upload(name, content):
    raw = content.encode() if isinstance(content, str) else content
    return {"relative_path": name, "content_b64": base64.b64encode(raw).decode()}


def office_upload(suffix):
    """Small real document payloads, independent of developer machine artifacts."""
    stream = io.BytesIO()
    if suffix == ".docx":
        from docx import Document

        document = Document()
        document.add_heading("Source lesson", level=1)
        document.add_paragraph("Confirm the issue before escalating.")
        document.save(stream)
    elif suffix == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["step", "action"])
        workbook.active.append([1, "Confirm the issue"])
        workbook.save(stream)
        workbook.close()
    elif suffix == ".pptx":
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("ppt/slides/slide1.xml",
                             '<p:sld xmlns:p="p" xmlns:a="a"><a:t>Confirm the issue</a:t></p:sld>')
    elif suffix == ".pdf":
        # Minimal valid text PDF; pypdf must recover real page text.
        objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
                   b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                   (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] "
                    b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"),
                   b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
        text = b"BT /F1 12 Tf 20 200 Td (Confirm the issue) Tj ET"
        objects.append(b"<< /Length " + str(len(text)).encode() + b" >>\nstream\n" + text + b"\nendstream")
        stream.write(b"%PDF-1.4\n")
        offsets = [0]
        for number, obj in enumerate(objects, start=1):
            offsets.append(stream.tell())
            stream.write(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
        xref = stream.tell()
        stream.write(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
        for offset in offsets[1:]:
            stream.write(f"{offset:010d} 00000 n \n".encode())
        stream.write(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    else:
        raise AssertionError(suffix)
    return stream.getvalue()
